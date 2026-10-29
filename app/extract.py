"""文件内容抽取：Office、PDF、文本；Mac 上用 textutil 兜底 .doc/.rtf；图片和扫描件走本地 OCR。"""
from __future__ import annotations

import re
import subprocess
import sys
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional

from . import standards as S

CJK = re.compile(r"[一-鿿]")
DATE_PATTERNS = [
    re.compile(r"(20\d{2})\s*[年\-/\.]\s*(\d{1,2})\s*[月\-/\.]\s*(\d{1,2})\s*日?"),
    re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)"),
]


def meaningful_len(s: str) -> int:
    """有意义的字数（汉字、字母、数字），用来判断 PDF 是不是扫描件。"""
    return len(re.findall(r"[一-鿿A-Za-z0-9]", s or ""))


def extract(path: Path, limit: int = 4000, ocr_enabled: bool = False, pdf_pages: int = 2) -> Dict:
    """返回 {text, cover_title, meta_title, kind, ocr...}；失败不抛异常。

    ocr_enabled 时：图片识别图中文字；PDF 没有文字层（扫描件）时渲染前几页识别。
    """
    ext = path.suffix.lower()
    out = {"text": "", "cover_title": "", "meta_title": "", "kind": ext.lstrip(".") or "file"}
    try:
        if ext == ".docx":
            _docx(path, out, limit)
        elif ext in (".xlsx", ".xlsm"):
            _xlsx(path, out, limit)
        elif ext == ".pptx":
            _pptx(path, out, limit)
        elif ext == ".pdf":
            _pdf(path, out, limit)
        elif ext in (".txt", ".md", ".csv", ".log", ".json", ".xml", ".html", ".htm", ".yaml", ".yml"):
            out["text"] = _read_text(path, limit)
        elif ext in (".doc", ".rtf", ".odt", ".wps", ".pages") and sys.platform == "darwin":
            out["text"] = _textutil(path, limit)
    except Exception as e:  # 解析失败只影响识别，不影响归档
        out["error"] = str(e)[:200]
    if ext in S.IMAGE_EXT or (ext == ".pdf" and meaningful_len(out["text"]) < 50):
        _ocr(path, ext, out, ocr_enabled, pdf_pages)
    out["text"] = _clean(out["text"])[:limit]
    if not out["cover_title"]:
        out["cover_title"] = _guess_title_from_text(out["text"])
    return out


def _ocr(path: Path, ext: str, out: Dict, enabled: bool, pdf_pages: int) -> None:
    """图片 / 扫描版 PDF 的文字识别。结果写进 out：ocr=True、ocr_ms；没做时 ocr_skip 写原因。"""
    from . import ocr

    is_pdf = ext == ".pdf"
    if not enabled:
        out["ocr_skip"] = "设置中已关闭图片文字识别"
        return
    if not ocr.ready() or (is_pdf and not ocr.pdf_render_installed()):
        out["ocr_skip"] = "未安装图片文字识别组件"
        return
    try:
        if is_pdf:
            text, n, sec = ocr.pdf_text(path, pdf_pages)
            out["ocr_pages"] = n
        else:
            text, size, sec = ocr.image_text(path)
            out["image_size"] = list(size)
    except Exception as e:
        out["ocr_error"] = str(e)[:160]
        return
    out["ocr"] = True
    out["ocr_ms"] = int(sec * 1000)
    if text:
        # 扫描件原有的少量文字层（如水印）放在识别结果后面
        out["text"] = text + ("\n" + out["text"] if out["text"].strip() else "")


def _clean(s: str) -> str:
    s = s.replace("　", " ").replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return s.strip()


def _docx(path: Path, out: Dict, limit: int) -> None:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts: List[str] = []
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        parts.append(t)
        style = (p.style.name or "").lower() if p.style is not None else ""
        if not out["cover_title"] and i < 40 and ("title" in style or "标题" in style or "heading" in style):
            if 2 <= len(t) <= 50:
                out["cover_title"] = t
        if sum(len(x) for x in parts) > limit:
            break
    if sum(len(x) for x in parts) < limit:
        for table in d.tables[:5]:
            for row in table.rows[:30]:
                cells = []
                for c in row.cells:
                    ct = c.text.strip()
                    if ct and ct not in cells:
                        cells.append(ct)
                if cells:
                    parts.append(" | ".join(cells))
    out["text"] = "\n".join(parts)
    try:
        out["meta_title"] = (d.core_properties.title or "").strip()
    except Exception:
        pass


def _xlsx(path: Path, out: Dict, limit: int) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: List[str] = []
    total = 0
    for ws in wb.worksheets[:4]:
        parts.append(f"[工作表] {ws.title}")
        for row in ws.iter_rows(max_row=80, values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                line = " | ".join(vals)
                parts.append(line)
                total += len(line)
                if not out["cover_title"] and 2 <= len(vals[0]) <= 50 and len(vals) <= 2:
                    out["cover_title"] = vals[0]
            if total > limit:
                break
        if total > limit:
            break
    wb.close()
    out["text"] = "\n".join(parts)


def _pptx(path: Path, out: Dict, limit: int) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: List[str] = []
    for idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                t = shape.text_frame.text.strip()
                parts.append(t)
                if idx == 0 and not out["cover_title"] and 2 <= len(t) <= 50:
                    out["cover_title"] = t.split("\n")[0]
        if sum(len(x) for x in parts) > limit:
            break
    out["text"] = "\n".join(parts)


def _pdf(path: Path, out: Dict, limit: int) -> None:
    from pypdf import PdfReader

    r = PdfReader(str(path))
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            return
    try:
        meta = r.metadata
        if meta and meta.title:
            out["meta_title"] = str(meta.title).strip()
    except Exception:
        pass
    parts: List[str] = []
    for page in r.pages[:4]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
        if sum(len(x) for x in parts) > limit:
            break
    out["text"] = "\n".join(parts)


def _read_text(path: Path, limit: int) -> str:
    raw = path.read_bytes()[: limit * 4]
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _textutil(path: Path, limit: int) -> str:
    try:
        res = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                             capture_output=True, timeout=30)
        return res.stdout.decode("utf-8", errors="ignore")[:limit]
    except Exception:
        return ""


def _guess_title_from_text(text: str) -> str:
    for line in text.split("\n")[:15]:
        t = line.strip(" |:：-")
        if 4 <= len(t) <= 40 and CJK.search(t) and not re.match(r"^[\d\s第章节一二三四五六七八九十、.]+$", t):
            return t
    return ""


def find_dates(*texts: str) -> List[str]:
    """抽取 2015 年至今（含 60 天以内的未来）的日期，按出现顺序返回 YYYY-MM-DD。"""
    today = date.today()
    found: List[str] = []
    for text in texts:
        if not text:
            continue
        for pat in DATE_PATTERNS:
            for m in pat.finditer(text):
                try:
                    d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                except ValueError:
                    continue
                if 2015 <= d.year and (d - today).days <= 60:
                    s = d.isoformat()
                    if s not in found:
                        found.append(s)
    return found


def pick_doc_date(name: str, text: str) -> Optional[str]:
    """优先文件名中的日期，其次正文里「日期：」后面的落款日期，再其次正文中最晚的日期。"""
    in_name = find_dates(name)
    if in_name:
        return in_name[-1]
    labeled = [find_dates(m.group(1)) for m in re.finditer(
        r"日\s*期\s*[:：]?\s*(20\d{2}\s*[年\-/\.]\s*\d{1,2}\s*[月\-/\.]\s*\d{1,2})", text[:3000])]
    labeled = [d[0] for d in labeled if d]
    if labeled:
        return labeled[-1]
    in_text = find_dates(text[:3000])
    if in_text:
        return max(in_text)
    return None
